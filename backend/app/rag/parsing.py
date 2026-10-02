"""Extract plain text from uploaded documents and split it into chunks."""
from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm", ".epub", ".csv", ".json", ".rtf"}
MAX_DOCUMENT_BYTES = 100 * 1024 * 1024
CHUNK_CHARS = 1_000
CHUNK_OVERLAP = 150


class ParseError(RuntimeError):
    pass


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "nav", "footer", "header", "svg"}

    def __init__(self) -> None:
        super().__init__()
        self._skip = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in self.SKIP:
            self._skip += 1
        elif tag in {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr", "section", "article"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    return _normalise("".join(parser.parts))


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ParseError("pypdf is not installed") from error
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as error:  # noqa: BLE001
            raise ParseError("PDF is password protected") from error
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - skip broken pages
            continue
    return "\n\n".join(pages)


def _docx(path: Path) -> str:
    try:
        import docx
    except ImportError as error:
        raise ParseError("python-docx is not installed") from error
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _epub(path: Path) -> str:
    try:
        import ebooklib
        from ebooklib import epub
    except ImportError as error:
        raise ParseError("ebooklib is not installed") from error
    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    parts = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        parts.append(html_to_text(item.get_content().decode("utf-8", errors="replace")))
    return "\n\n".join(parts)


def _text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: Path, name: str) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        text = _pdf(path)
    elif suffix == ".docx":
        text = _docx(path)
    elif suffix == ".epub":
        text = _epub(path)
    elif suffix in {".html", ".htm"}:
        text = html_to_text(_text(path))
    elif suffix in ALLOWED_DOCUMENT_EXTENSIONS:
        text = _text(path)
    else:
        raise ParseError(f"Unsupported document type {suffix or '(none)'}")
    text = _normalise(text)
    if len(text) < 20:
        raise ParseError("No readable text found in the document (scanned PDFs need OCR first)")
    return text


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware sliding window so chunks end on sentence boundaries where possible."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(_split_long(paragraph, size, overlap))
            continue
        if len(current) + len(paragraph) + 2 > size and current:
            chunks.append(current)
            current = current[-overlap:].strip() + "\n" if overlap else ""
        current = f"{current}{paragraph}\n\n" if current else paragraph + "\n\n"
    if current.strip():
        chunks.append(current)
    return [c.strip() for c in chunks if len(c.strip()) >= 40]


def _split_long(paragraph: str, size: int, overlap: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+", paragraph)
    pieces: list[str] = []
    current = ""
    for sentence in sentences:
        while len(sentence) > size:
            pieces.append(sentence[:size])
            sentence = sentence[size - overlap :]
        if len(current) + len(sentence) + 1 > size and current:
            pieces.append(current)
            current = current[-overlap:] if overlap else ""
        current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return pieces
